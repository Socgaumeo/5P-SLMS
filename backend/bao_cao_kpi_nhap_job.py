"""Báo cáo KPI NHẬP JOB theo bước (file V1). Chạy:  python bao_cao_kpi_nhap_job.py <file.xlsx> [job_id_tu]"""
import sys, datetime as dt, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from app.db.supabase_client import get_supabase
from app.services import kpi_buoc as K

out = sys.argv[1]
tu = int(sys.argv[2]) if len(sys.argv) > 2 else 0
c = get_supabase()
now = dt.datetime.now(K.VN)
users = {u["user_id"]: u["full_name"] for u in c.table("users").select("user_id,full_name").execute().data}
import os as _os
_q = c.table("jobs").select("job_id").gte("job_id", tu).neq("status_code", "CANCELLED")
_q = _q.gte("created_at", _os.environ.get("KPI_AP_DUNG_TU", "2026-10-12T21:00:00+07:00"))
jobs = _q.order("job_id").execute().data
done = {(d["svc_id"], d["buoc"]): d for d in c.table("job_step_done").select("*").gte("job_id", tu).execute().data}

rows = []
for j in jobs:
    for r in K.danh_gia_job(c, j["job_id"]):
        d = done.get((r["svc_id"], r["buoc"]))
        du_luc = K._ts(d["done_at"]).astimezone(K.VN) if d else None
        nguoi = d["actor"] if d else ""
        uid = nguoi.split("|")[0] if nguoi else ""
        tu_dong = uid == "auto"
        kenh = ("Tự động" if tu_dong else
                {"sen": "Sen nhập hộ", "web": "Tự nhập web"}.get(nguoi.split("|")[1] if "|" in nguoi else "",
                                                                "Không rõ" if d else ""))
        nguoi_ten = (f"Tự động ({nguoi.split('|')[1]})" if tu_dong else
                     users.get(int(uid)) if uid.isdigit() else ("(không rõ)" if d else ""))
        h = r["han"]
        if r.get("loi"):
            l = r["loi"].lower()
            kq = ("Lỗi: thiếu kỳ thanh toán" if "kỳ thanh toán" in l else
                  "Lỗi: không cập nhật lịch" if "lịch dự kiến" in l else
                  "Lỗi: bỏ qua bước" if "bỏ qua bước" in l else "Lỗi: " + l)
        elif du_luc:
            kq = "Chưa đến bước" if not h else ("Đạt" if du_luc <= h else "Trễ")
        else:
            kq = "Chưa đến bước" if not h else ("Quá hạn, chưa đủ" if now > h else "Trong hạn, chưa đủ")
        tre = round((du_luc - h).total_seconds() / 3600, 1) if (du_luc and h and du_luc > h) else (
            round((now - h).total_seconds() / 3600, 1) if (not du_luc and h and now > h) else None)
        rows.append({
            "Job": r["job_no"], "Người tạo job": users.get(r["created_by"], ""), "Dịch vụ": r["svc_type"],
            "Bước": f'{r["buoc"]}. {r["ten_buoc"]}', "Trạng thái V1": r["trang_thai"], "KPI": r["kpi"],
            "Mốc T (thao tác)": r["ten_moc"], "Giờ T": r["T"].astimezone(K.VN).strftime("%d/%m %H:%M") if r["T"] else "",
            "Hạn": h.strftime("%d/%m %H:%M") if h else "", "Nhập đủ lúc": du_luc.strftime("%d/%m %H:%M") if du_luc else "",
            "Người nhập": nguoi_ten, "Kênh": kenh, "Người chịu KPI": "(tự động — không tính KPI)" if tu_dong else (nguoi_ten if uid.isdigit() else users.get(r["created_by"], "") or "(không rõ)"), "Kết quả": kq, "Trễ (giờ)": tre,
            "Lịch dự kiến": r["lich_du_kien"].strftime("%d/%m %H:%M") if r.get("lich_du_kien") else "", "Hạn cập nhật lịch (+6h)": r["han_cap_nhat_lich"].strftime("%d/%m %H:%M") if r.get("han_cap_nhat_lich") else "", "Lỗi": r.get("loi") or "", "Còn thiếu ô": ", ".join(K.NHAN.get(x, x) for x in r["thieu_o"]),
            "Còn thiếu chứng từ": ", ".join(K.DOC_NHAN.get(x, x) for x in r["thieu_ct"]),
        })

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); F = PatternFill("solid", fgColor="1F3864")
MAU = {"Lỗi: thiếu kỳ thanh toán": "F4B084", "Lỗi: không cập nhật lịch": "F4B084", "Lỗi: bỏ qua bước": "F4B084", "Đạt": "C6EFCE", "Trễ": "FFC7CE", "Quá hạn, chưa đủ": "FFC7CE", "Trong hạn, chưa đủ": "FFEB9C", "Chưa đến bước": "EDEDED"}


def sheet(ws, header, data, color_col=None):
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = H, F
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for d in data:
        ws.append([d.get(h) for h in header])
        if color_col and d.get(color_col) in MAU:
            ws.cell(ws.max_row, header.index(color_col) + 1).fill = PatternFill("solid", fgColor=MAU[d[color_col]])
    for i, h in enumerate(header, 1):
        w = max([len(str(h))] + [len(str(d.get(h) or "")) for d in data]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 8), 50)
    ws.freeze_panes = "A2"


def tong(key):
    g = collections.OrderedDict()
    for r in rows:
        if key in ("Người nhập", "Người chịu KPI") and r.get("Kênh") == "Tự động":
            continue          # Sen tự cập nhật từ mail/tracking: không tính KPI cho nhân viên
        k = r[key] or "(chưa nhập)"
        x = g.setdefault(k, collections.Counter())
        x[r["Kết quả"]] += 1
    out = []
    for k, x in g.items():
        loi = sum(v for kk, v in x.items() if kk.startswith("Lỗi"))
        mau = x["Đạt"] + x["Trễ"] + x["Quá hạn, chưa đủ"] + loi     # "Chưa đến bước" không tính
        out.append({key: k, "Đạt": x["Đạt"], "Trễ": x["Trễ"], "Quá hạn, chưa đủ": x["Quá hạn, chưa đủ"], "Lỗi": loi,
                    "Trong hạn, chưa đủ": x["Trong hạn, chưa đủ"], "Chưa đến bước": x["Chưa đến bước"],
                    "Tỷ lệ đúng hạn": f"{x['Đạt'] / mau * 100:.0f}%" if mau else "—"})
    return out


ws = wb.active; ws.title = "Tổng hợp"
ws.append([f"BÁO CÁO KPI NHẬP JOB — bản test, lập {now:%d/%m/%Y %H:%M}"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Cách tính: mọi mốc giờ = thời điểm nhân viên/Sen thao tác trên hệ thống (không nhập tay giờ). Mốc T = thao tác đầu tiên báo việc xảy ra "
           "(tạo job, gán xe, báo đã giao/tải BBGH, nhập số TK, nhập vận đơn). Sen nhập hộ: tính cho nhân viên gửi, giờ = lúc nhân viên gửi tin/file. "
           "Bước đạt khi nhập ĐỦ ô + chứng từ bắt buộc trong hạn file V1. Tỷ lệ đúng hạn = Đạt / (Đạt + Trễ + Quá hạn chưa đủ + Lỗi); \"Chưa đến bước\" (việc chưa xảy ra) không tính. Bước có lịch dự kiến: quá lịch + 6h mà chưa có thao tác thực tế và chưa đổi lịch = lỗi."])
ws.append([])
cnt = collections.Counter(r["Kết quả"] for r in rows)
for k in ["Đạt", "Trễ", "Quá hạn, chưa đủ", "Trong hạn, chưa đủ", "Chưa đến bước", "Lỗi: thiếu kỳ thanh toán", "Lỗi: không cập nhật lịch", "Lỗi: bỏ qua bước"]:
    ws.append([k, cnt[k]])
    if k in MAU:
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=MAU[k])
ws.append(["Tổng bước", len(rows)]); ws.append(["Số job", len({r['Job'] for r in rows})])
ws.column_dimensions["A"].width = 24
sheet(wb.create_sheet("Theo nhân viên"), ["Người chịu KPI", "Đạt", "Trễ", "Quá hạn, chưa đủ", "Lỗi", "Trong hạn, chưa đủ", "Chưa đến bước", "Tỷ lệ đúng hạn"], tong("Người chịu KPI"))
sheet(wb.create_sheet("Theo bước"), ["Bước", "Đạt", "Trễ", "Quá hạn, chưa đủ", "Lỗi", "Trong hạn, chưa đủ", "Chưa đến bước", "Tỷ lệ đúng hạn"], tong("Bước"))
sheet(wb.create_sheet("Chi tiết"), list(rows[0].keys()) if rows else ["Job"], rows, "Kết quả")
# Lô cần Sen kiểm lại lịch với nhân viên / tracking hãng bay-tàu / BBGH
kiem = [{k: r[k] for k in ("Job", "Người chịu KPI", "Dịch vụ", "Bước", "Lịch dự kiến", "Hạn cập nhật lịch (+6h)", "Lỗi")}
        for r in rows if r["Kết quả"] == "Lỗi: không cập nhật lịch"]
sheet(wb.create_sheet("Cần kiểm lại lịch"), ["Job", "Người chịu KPI", "Dịch vụ", "Bước", "Lịch dự kiến", "Hạn cập nhật lịch (+6h)", "Lỗi"], kiem)
wb.save(out)
print("saved", out, len(rows), dict(cnt))
